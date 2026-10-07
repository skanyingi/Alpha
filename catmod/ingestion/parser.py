"""Parse messy cedant/broker bordereaux (CSV, XLSX) and PDF claim notices."""

from __future__ import annotations

import base64
import csv
import io
import re
from typing import Any

from catmod.schemas import RawClaim

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "asset_id": ("asset_id", "assetid", "location_id", "loc_id", "risk_id", "id"),
    "policy_id": ("policy_id", "policyid", "policy", "pol_id", "contract_id"),
    "latitude": ("latitude", "lat", "y", "coord_lat"),
    "longitude": ("longitude", "lon", "lng", "long", "x", "coord_lon"),
    "elevation": ("elevation", "elev", "alt", "altitude", "z"),
    "occupancy_raw": (
        "occupancy",
        "occupancy_raw",
        "occ",
        "occupancy_type",
        "construction",
        "use",
        "bldg_occ",
    ),
    "tiv": ("tiv", "total_insured_value", "sum_insured", "si", "insured_value", "value"),
    "ground_up_loss": (
        "ground_up_loss",
        "gul",
        "loss",
        "claim",
        "claim_amount",
        "gross_loss",
        "damage",
    ),
    "deductible": ("deductible", "ded", "excess", "franchise"),
    "policy_limit": ("policy_limit", "limit", "pol_limit", "occurrence_limit"),
    "coinsurance": ("coinsurance", "coins", "co_insurance", "insured_share"),
    "loss_date": ("loss_date", "date", "event_date", "doi", "date_of_loss"),
}


def _norm_header(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def _bind_headers(headers: list[str]) -> dict[str, str]:
    normalized = [_norm_header(h) for h in headers]
    bound: dict[str, str] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for header, norm in zip(headers, normalized):
            if norm in aliases and field not in bound:
                bound[field] = header
                break
    return bound


def _num(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    text = str(value).strip()
    if text == "" or text.lower() in {"na", "n/a", "null", "none"}:
        return default
    text = text.replace(",", "").replace("$", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return default


def _str(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def rows_to_claims(headers: list[str], rows: list[dict[str, Any]]) -> list[RawClaim]:
    bound = _bind_headers(headers)
    claims: list[RawClaim] = []
    for i, row in enumerate(rows, start=1):
        def pick(field: str, default: Any = "") -> Any:
            src = bound.get(field)
            if src is None:
                return default
            return row.get(src, default)

        asset = _str(pick("asset_id"), f"ROW-{i:05d}")
        extras = {
            k: v
            for k, v in row.items()
            if k not in bound.values()
        }
        claims.append(
            RawClaim(
                asset_id=asset or f"ROW-{i:05d}",
                policy_id=_str(pick("policy_id")),
                latitude=_num(pick("latitude")),
                longitude=_num(pick("longitude")),
                elevation=_num(pick("elevation")),
                occupancy_raw=_str(pick("occupancy_raw")),
                tiv=_num(pick("tiv")),
                ground_up_loss=_num(pick("ground_up_loss")),
                deductible=_num(pick("deductible")),
                policy_limit=_num(pick("policy_limit")),
                coinsurance=_num(pick("coinsurance"), 1.0) or 1.0,
                loss_date=_str(pick("loss_date")),
                extras=extras,
            )
        )
    return claims


def parse_csv_text(text: str) -> list[RawClaim]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = list(reader.fieldnames or [])
    rows = [dict(r) for r in reader]
    return rows_to_claims(headers, rows)


def parse_xlsx_bytes(blob: bytes) -> list[RawClaim]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(header_row)]
    rows: list[dict[str, Any]] = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append({h: v for h, v in zip(headers, values)})
    return rows_to_claims(headers, rows)


def parse_pdf_bytes(blob: bytes) -> list[RawClaim]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return parse_claim_narrative(text)


_MONEY_RE = re.compile(r"\$?\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)")
_COORD_RE = re.compile(
    r"(-?\d{1,3}\.\d+)\s*[,;\s]\s*(-?\d{1,3}\.\d+)",
)


def parse_claim_narrative(text: str) -> list[RawClaim]:
    """Best-effort extraction from a PDF/email claim notification."""
    lat, lon = 0.0, 0.0
    coord = _COORD_RE.search(text)
    if coord:
        lat, lon = float(coord.group(1)), float(coord.group(2))
    amounts = [_num(m.group(1)) for m in _MONEY_RE.finditer(text)]
    gul = max(amounts) if amounts else 0.0
    occupancy = ""
    occ_match = re.search(
        r"(warehouse|whse|office|dwelling|residential|retail|industrial|hotel|storage|shed)",
        text,
        re.I,
    )
    if occ_match:
        occupancy = occ_match.group(1)
    return [
        RawClaim(
            asset_id="PDF-CLAIM-001",
            occupancy_raw=occupancy or "unstructured_pdf",
            latitude=lat,
            longitude=lon,
            ground_up_loss=gul,
            extras={"narrative": text[:4000]},
        )
    ]


def parse_attachment(
    *,
    filename: str,
    text: str | None = None,
    blob: bytes | None = None,
    content_type: str | None = None,
) -> list[RawClaim]:
    name = filename.lower()
    ctype = (content_type or "").lower()
    if blob is None and text is not None and (name.endswith(".xlsx") or "spreadsheet" in ctype):
        blob = text.encode("latin-1", errors="ignore")
    if name.endswith(".xlsx") or "spreadsheet" in ctype:
        if blob is None:
            raise ValueError("XLSX attachment requires binary data")
        return parse_xlsx_bytes(blob)
    if name.endswith(".pdf") or "pdf" in ctype:
        if blob is None:
            raise ValueError("PDF attachment requires binary data")
        return parse_pdf_bytes(blob)
    if text is None and blob is not None:
        text = blob.decode("utf-8", errors="replace")
    if text is None:
        raise ValueError("CSV attachment is empty")
    return parse_csv_text(text)


def decode_payload_bytes(data_base64: str | None) -> bytes | None:
    if not data_base64:
        return None
    return base64.b64decode(data_base64)
